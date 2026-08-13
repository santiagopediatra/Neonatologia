#!/usr/bin/env python3
"""Corrige solo pendientes de la auditoría crítica del DOCX maestro."""

from __future__ import annotations

import csv
import re
import shutil
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "manuscrito/Articulo_fetos_MAESTRO.docx"
BACKUP = ROOT / "manuscrito/Articulo_fetos_MAESTRO_BACKUP.docx"
OUTPUT = ROOT / "manuscrito/Articulo_fetos_MAESTRO_FINAL.docx"
MISSING_OUT = ROOT / "AUDITORIA_DATOS_FALTANTES.csv"


def replace_text(paragraph, old: str, new: str) -> None:
    """Reemplaza texto conservando los formatos de los runs cuando es posible."""
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return
    if old in paragraph.text:
        # Los párrafos afectados son prosa normal, sin énfasis interno relevante.
        paragraph.text = paragraph.text.replace(old, new)


def append_sentence(paragraph, sentence: str) -> None:
    if sentence not in paragraph.text:
        paragraph.add_run((" " if paragraph.text else "") + sentence)


def insert_after(paragraph, text: str, style: str | None = None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        try:
            new_para.style = style
        except KeyError:
            pass
    new_para.add_run(text)
    return new_para


def build_missing_audit() -> None:
    base_path = ROOT / "datos/procesados/base_estudio_muerte_fetal.csv"
    df = pd.read_csv(base_path, low_memory=False)
    model_raw = {
        "MUERTE_FETAL_BINARIA", "EDAD_CATEGORIA_ANALITICA", "PAREJA_ESTABLE",
        "SEGURO", "INSTRUCCION", "ACOMPANADA", "RIESGO_AL_INGRESO",
        "NO_CONTROLES_NUM", "IVU", "RPM", "SESIONES_NUM", "HEMORRAGIA",
        "PRESENTACION_FETAL", "HIPERTENSIVOS", "DESPRENDIMIENTO_PREMATURO_DE_PLACENTA",
        "EDAD_GESTACIONAL_ANALITICA",
    }
    rows = []
    total = len(df)
    for col in df.columns:
        available = int(df[col].notna().sum())
        missing = total - available
        rows.append({
            "variable": col,
            "N total": total,
            "N disponible": available,
            "N faltante": missing,
            "porcentaje faltante": round(100 * missing / total, 2),
            "participa en modelo sí/no": "sí" if col in model_raw else "no",
            "población de referencia": "base total",
        })

    modeled_path = ROOT / "resultados/tablas/04_faltantes_variables_modeladas.csv"
    modeled = pd.read_csv(modeled_path)
    analytic_total = int((modeled["N_VALIDO"] + modeled["N_FALTANTE"]).max())
    for _, r in modeled.iterrows():
        rows.append({
            "variable": r["VARIABLE"],
            "N total": analytic_total,
            "N disponible": int(r["N_VALIDO"]),
            "N faltante": int(r["N_FALTANTE"]),
            "porcentaje faltante": round(float(r["PCT_FALTANTE"]), 2),
            "participa en modelo sí/no": "sí",
            "población de referencia": "registros con desenlace conocido",
        })
    pd.DataFrame(rows).to_csv(MISSING_OUT, index=False, encoding="utf-8")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not BACKUP.exists():
        shutil.copy2(SOURCE, BACKUP)

    build_missing_audit()
    doc = Document(SOURCE)
    original_tables = [[cell.text for row in table.rows for cell in row.cells] for table in doc.tables]
    original_shapes = len(doc.inline_shapes)

    # Referencias: solo cuatro marcadores comprobados.
    citations = {
        "episodio obstétrico [CITA]": "episodio obstétrico [1,2]",
        "desenlaces fetales adversos [CITA]": "desenlaces fetales adversos [3]",
        "vía causal [CITA]": "vía causal [4]",
        "selección de pacientes [CITA]": "selección de pacientes [4]",
    }
    for p in doc.paragraphs:
        for old, new in citations.items():
            replace_text(p, old, new)

    # Adherencia: precisión solicitada sin recalcular ni modificar resultados.
    for p in doc.paragraphs:
        if p.text.startswith("Para el análisis adicional de adherencia"):
            append_sentence(
                p,
                "Como no hubo registros analíticos con edad gestacional menor de 20 semanas, en la población analítica real se aplicaron principalmente los puntos 20–<28=3, 28–<36=6 y ≥36=8 controles esperados.",
            )

    risk_sentence = (
        "Las categorías de riesgo al ingreso son evaluaciones clínicas realizadas al momento de la atención; "
        "se presentan para caracterizar la gravedad clínica asociada con los casos y no como exposiciones "
        "etiológicas independientes ni como estimaciones causales."
    )
    for p in doc.paragraphs:
        if p.text.startswith("Las variables maternas fueron"):
            append_sentence(p, risk_sentence)
        elif p.text.startswith("Resultados. Se registraron"):
            replace_text(
                p,
                "El riesgo alto e inminente al ingreso presentó OR posteriores de 1,60 y 1,81.",
                "El riesgo alto e inminente al ingreso, interpretado como marcador clínico y no como exposición etiológica, presentó OR posteriores de 1,60 y 1,81.",
            )
        elif p.text.startswith("Conclusiones. DPP"):
            replace_text(
                p,
                "DPP y mayor gravedad al ingreso mostraron asociaciones positivas claras",
                "DPP mostró una asociación positiva clara; las categorías de mayor gravedad al ingreso también se asociaron positivamente y se interpretaron como marcadores clínicos, no como exposiciones etiológicas",
            )
        elif p.text.startswith("Riesgo alto e inminente se asociaron"):
            append_sentence(p, risk_sentence)
        elif p.text.startswith("En estos registros obstétricos, el DPP"):
            replace_text(
                p,
                "deben entenderse principalmente como marcadores de gravedad clínica.",
                "deben entenderse como evaluaciones clínicas al momento de la atención y marcadores de gravedad, no como exposiciones etiológicas independientes ni estimaciones causales.",
            )

    # Sustituye la nota bibliográfica vacía por fuentes verificadas.
    ref_note = next(p for p in doc.paragraphs if p.text.startswith("No se incorporaron referencias bibliográficas"))
    ref_note.text = (
        "1. World Health Organization. Stillbirth. Disponible en: "
        "https://www.who.int/health-topics/stillbirth (consulta: 12 de agosto de 2026)."
    )
    current = ref_note
    refs = [
        "2. Aminu M, Bar-Zeev S, van den Broek N. Cause of and factors associated with stillbirth: a systematic review of classification systems. Acta Obstet Gynecol Scand. 2017;96(5):519–528. doi:10.1111/aogs.13126.",
        "3. Downes KL, Grantz KL, Shenassa ED. Maternal, labor, delivery, and perinatal outcomes associated with placental abruption: a systematic review. Am J Perinatol. 2017;34(10):935–957. doi:10.1055/s-0037-1599149.",
        "4. Vandenbroucke JP, von Elm E, Altman DG, et al. Strengthening the Reporting of Observational Studies in Epidemiology (STROBE): explanation and elaboration. Epidemiology. 2007;18(6):805–835. doi:10.1097/EDE.0b013e3181577511.",
    ]
    for ref in refs:
        current = insert_after(current, ref, ref_note.style.name)

    doc.save(OUTPUT)

    # Garantías de no alteración de tablas/figuras y de eliminación de citas pendientes.
    check = Document(OUTPUT)
    final_tables = [[cell.text for row in table.rows for cell in row.cells] for table in check.tables]
    assert final_tables == original_tables, "Se alteró una tabla; operación abortada"
    assert len(check.inline_shapes) == original_shapes, "Cambió el número de figuras"
    full_text = "\n".join(p.text for p in check.paragraphs)
    assert "[CITA]" not in full_text
    assert "[CITAR]" not in full_text and "[REFERENCIA]" not in full_text

    # Tabla 1: cada valor reproducible debe estar presente, con coma decimal del Word.
    with (ROOT / "tabla1_corregida.csv").open(encoding="utf-8-sig") as fh:
        expected = list(csv.DictReader(fh))
    table1_text = "\n".join(final_tables[0])
    missing_values = []
    for row in expected:
        value = row["valor_formateado"]
        value_comma = value.replace(".", ",")
        if value not in table1_text and value_comma not in table1_text:
            missing_values.append(value)
    assert not missing_values, f"Valores ausentes en Tabla 1: {missing_values[:5]}"

    print(f"Creado: {OUTPUT}")
    print(f"Tablas preservadas: {len(final_tables)}; figuras preservadas: {len(check.inline_shapes)}")
    print(f"Marcadores [COMPLETAR] autorales conservados: {full_text.count('[COMPLETAR]')}")


if __name__ == "__main__":
    main()
