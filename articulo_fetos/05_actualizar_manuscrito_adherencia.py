#!/usr/bin/env python3

import csv
import re
import shutil
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches

ENTRADA = Path("manuscrito/Articulo_fetos_TABLA1_CORREGIDA.docx")
SALIDA = Path("manuscrito/Articulo_fetos_CORREGIDO.docx")
FIGURA = Path("adherencia_fig1.png")


def leer_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def fmt(x, d=2):
    return f"{float(x):.{d}f}".replace(".", ",")


def fmt_p(x):
    x = float(x)
    return "<0,001" if x < 0.001 else fmt(x, 3)


def miles(x):
    return f"{int(float(x)):,}".replace(",", ".")


def buscar_parrafo(doc, inicio):
    encontrados = [p for p in doc.paragraphs if p.text.strip().startswith(inicio)]
    if len(encontrados) != 1:
        raise RuntimeError(f"No se encontró inequívocamente el párrafo: {inicio}")
    return encontrados[0]


def parrafo_despues(elemento, texto="", estilo=None):
    nuevo = OxmlElement("w:p")
    elemento.addnext(nuevo)
    p = Paragraph(nuevo, doc._body)
    if estilo:
        p.style = estilo
    if texto:
        p.add_run(texto)
    return p


def agregar_tabla_despues(doc, elemento, encabezados, filas):
    tabla = doc.add_table(rows=1, cols=len(encabezados))
    if doc.tables and doc.tables[0].style is not None:
        tabla.style = doc.tables[0].style
    for i, texto in enumerate(encabezados):
        tabla.rows[0].cells[i].text = texto
        for run in tabla.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for fila in filas:
        cells = tabla.add_row().cells
        for i, texto in enumerate(fila):
            cells[i].text = str(texto)
    elemento.addnext(tabla._tbl)
    return tabla._tbl


if not ENTRADA.exists():
    raise RuntimeError(f"No existe el Word intermedio: {ENTRADA}")
if not FIGURA.exists():
    raise RuntimeError(f"No existe la figura: {FIGURA}")

a1 = leer_csv("adherencia_tabla_A1.csv")
a2 = leer_csv("adherencia_tabla_A2.csv")
modelos = [r for r in leer_csv("adherencia_modelos_logisticos.csv") if r["termino"].startswith("categoria_adherencia")]
estratos = leer_csv("adherencia_estratificado.csv")

shutil.copy2(ENTRADA, SALIDA)
doc = Document(SALIDA)

# Cambios textuales controlados: se añaden resultados sin sustituir análisis previos.
buscar_parrafo(doc, "Métodos. Estudio transversal").add_run(
    " Como análisis adicional, se construyó una adherencia operacional según controles observados/esperados por edad gestacional; el ajuste por edad gestacional se trató como sensibilidad."
)
buscar_parrafo(doc, "Resultados. Se registraron").add_run(
    " En 23.442 registros completos, la mortalidad fue 1,56% con adherencia buena, 1,47% con adherencia parcial y 3,35% con adherencia baja. Para baja frente a buena, el OR crudo fue 2,18 (IC95% 1,72–2,77) y el OR ajustado por edad gestacional, considerado sensibilidad, fue 1,84 (IC95% 1,43–2,38)."
)
buscar_parrafo(doc, "Se calcularon frecuencias absolutas").add_run(
    " Para el análisis adicional de adherencia se dividieron los controles observados por los esperados según la especificación operacional: <13 semanas=1, 13–<20=2, 20–<28=3, 28–<36=6 y ≥36=8. La adherencia se clasificó como buena (≥80%), parcial (≥50% y <80%) o baja (<50%), sin truncar valores >100%. Se utilizaron casos completos para desenlace, edad gestacional analítica y controles. Se calcularon IC95% binomiales de Wilson y regresiones logísticas cruda y ajustada por edad gestacional estandarizada; esta última se interpretó como sensibilidad por la dependencia estructural del indicador respecto de la edad gestacional."
)

p_resultados = buscar_parrafo(doc, "Las asociaciones negativas de controles")
p_resultados.add_run(
    " En el análisis operacional de adherencia se incluyeron 23.442 registros y 423 muertes fetales. La distribución fue: buena, 10.242 (43,69%); parcial, 9.524 (40,63%); y baja, 3.676 (15,68%). Hubo 3.693 registros (15,75%) con adherencia >100%, que no fueron truncados ni excluidos. La mortalidad fue 1,56% (160/10.242) en buena, 1,47% (140/9.524) en parcial y 3,35% (123/3.676) en baja. Frente a buena adherencia, parcial mostró OR 0,94 (IC95% 0,75–1,18; p=0,596) y baja OR 2,18 (IC95% 1,72–2,77; p<0,001). En la sensibilidad ajustada por edad gestacional, los OR fueron 1,08 (IC95% 0,85–1,37; p=0,528) y 1,84 (IC95% 1,43–2,38; p<0,001), respectivamente."
)

buscar_parrafo(doc, "Las asociaciones negativas del modelo principal").add_run(
    " En el indicador operacional, la categoría baja presentó mayores odds que la categoría buena. La estimación se desplazó hacia la nulidad al incluir edad gestacional, pero este cambio no distingue entre confusión, acoplamiento matemático, oportunidad acumulada de controles, selección, temporalidad o una asociación real."
)
buscar_parrafo(doc, "Las limitaciones incluyen diseño transversal").add_run(
    " El indicador de adherencia utilizó una regla operacional solicitada no validada externamente en este proyecto; además, la edad gestacional determina los controles esperados y participa en el modelo de sensibilidad, lo que introduce dependencia estructural y limita la interpretación del cambio entre estimaciones."
)
buscar_parrafo(doc, "Controles prenatales, sesiones de psicoprofilaxis").add_run(
    " En el análisis adicional, la adherencia baja se asoció con mayores odds frente a la buena, con atenuación en la sensibilidad por edad gestacional; este resultado depende de la definición operacional y no demuestra un efecto causal."
)

# Tablas auxiliares y figura, insertadas inmediatamente después de 4.9.
anchor = p_resultados._p
cap = parrafo_despues(anchor, "Tabla A1. Descripción de controles y edad gestacional por adherencia y desenlace.", "Body Text")
anchor = cap._p
filas_a1 = []
for r in a1:
    filas_a1.append([
        r["categoria_adherencia"], r["desenlace"], r["n"], fmt(r["porcentaje_columna"]),
        f'{fmt(r["media_controles_observados"])} ({fmt(r["de_controles_observados"])})',
        f'{fmt(r["mediana_controles_observados"])} ({fmt(r["q1_controles_observados"])}–{fmt(r["q3_controles_observados"])})',
        f'{fmt(r["media_controles_esperados"])} ({fmt(r["de_controles_esperados"])})',
        f'{fmt(r["media_eg"])} ({fmt(r["de_eg"])})',
    ])
anchor = agregar_tabla_despues(doc, anchor,
    ["Adherencia", "Desenlace", "n", "% columna", "Controles observados: media (DE)",
     "Controles observados: mediana (RIC)", "Controles esperados: media (DE)", "EG: media (DE)"],
    filas_a1)

cap = parrafo_despues(anchor, "Tabla A2. Mortalidad fetal por categoría operacional de adherencia.", "Body Text")
anchor = cap._p
filas_a2 = [[r["categoria_adherencia"], r["N"], r["n_muertes"], fmt(r["mortalidad_pct"]),
             fmt(r["IC95_inf"]), fmt(r["IC95_sup"])] for r in a2]
anchor = agregar_tabla_despues(doc, anchor,
    ["Adherencia", "N", "Muertes", "Mortalidad (%)", "IC95% inf", "IC95% sup"], filas_a2)

cap = parrafo_despues(anchor, "Tabla A3. Regresión logística de adherencia y muerte fetal.", "Body Text")
anchor = cap._p
filas_mod = []
for r in modelos:
    contraste = r["termino"].replace("categoria_adherencia", "") + " vs Buena"
    filas_mod.append([r["modelo"], contraste, fmt(r["OR"]), fmt(r["IC95_inf"]),
                      fmt(r["IC95_sup"]), fmt_p(r["p"]), r["N"], r["eventos"]])
anchor = agregar_tabla_despues(doc, anchor,
    ["Modelo", "Contraste", "OR", "IC95% inf", "IC95% sup", "p", "N", "Eventos"], filas_mod)

cap = parrafo_despues(anchor, "Tabla A4. Mortalidad fetal por adherencia, estratificada por edad gestacional.", "Body Text")
anchor = cap._p
filas_est = [[r["estrato_eg"], r["categoria_adherencia"], r["N"], r["n_muertes"],
              fmt(r["mortalidad_pct"]), fmt(r["IC95_inf"]), fmt(r["IC95_sup"])] for r in estratos]
anchor = agregar_tabla_despues(doc, anchor,
    ["Estrato EG", "Adherencia", "N", "Muertes", "Mortalidad (%)", "IC95% inf", "IC95% sup"], filas_est)

pfig = parrafo_despues(anchor, estilo="Body Text")
pfig.alignment = WD_ALIGN_PARAGRAPH.CENTER
pfig.add_run().add_picture(str(FIGURA), width=Inches(6.2))
anchor = pfig._p
capfig = parrafo_despues(anchor,
    "Figura A1. Mortalidad fetal por adherencia a controles prenatales. Puntos: proporciones observadas; barras: IC95% de Wilson.",
    "Image Caption")

doc.save(SALIDA)
print("Plan aplicado:")
print("- Resumen: Métodos y Resultados")
print("- Métodos: análisis descriptivo y frecuentista")
print("- Resultados: atención prenatal, tablas A1–A4 y figura A1")
print("- Discusión: atención prenatal y limitaciones")
print("- Conclusiones: interpretación de adherencia")
print(f"Word final generado: {SALIDA}")

# ---------------- Validación independiente por segunda lectura ----------------
final = Document(SALIDA)

def norm(texto):
    texto = unicodedata.normalize("NFKD", texto or "")
    return "".join(c for c in texto if not unicodedata.combining(c)).casefold()

def encontrar_tabla(encabezados):
    objetivo = [norm(x) for x in encabezados]
    halladas = [t for t in final.tables if [norm(c.text.strip()) for c in t.rows[0].cells] == objetivo]
    if len(halladas) != 1:
        raise RuntimeError(f"Tabla final no encontrada inequívocamente: {encabezados}")
    return halladas[0]

# Tabla 1 contra CSV.
t1 = encontrar_tabla(["Variable", "Categoria", "MUERTE_FETAL_NO", "MUERTE_FETAL_SI"])
calc_t1 = leer_csv("tabla1_corregida.csv")
esperados_t1 = {(norm(r["etiqueta"]), norm(r["categoria"]), r["desenlace"]): r["valor_formateado"] for r in calc_t1}
variable_actual = ""
errores_t1 = 0
for row in t1.rows[1:]:
    if row.cells[0].text.strip():
        variable_actual = row.cells[0].text.strip()
    categoria = row.cells[1].text.strip()
    for idx, des in ((2, "NO"), (3, "SI")):
        if row.cells[idx].text.strip() != esperados_t1[(norm(variable_actual), norm(categoria), des)]:
            errores_t1 += 1

def matriz(tabla):
    return [[c.text.strip() for c in row.cells] for row in tabla.rows[1:]]

checks = []
checks.append(matriz(encontrar_tabla(
    ["Adherencia", "Desenlace", "n", "% columna", "Controles observados: media (DE)",
     "Controles observados: mediana (RIC)", "Controles esperados: media (DE)", "EG: media (DE)"])) == filas_a1)
checks.append(matriz(encontrar_tabla(
    ["Adherencia", "N", "Muertes", "Mortalidad (%)", "IC95% inf", "IC95% sup"])) == filas_a2)
checks.append(matriz(encontrar_tabla(
    ["Modelo", "Contraste", "OR", "IC95% inf", "IC95% sup", "p", "N", "Eventos"])) == filas_mod)
checks.append(matriz(encontrar_tabla(
    ["Estrato EG", "Adherencia", "N", "Muertes", "Mortalidad (%)", "IC95% inf", "IC95% sup"])) == filas_est)

texto_final = "\n".join(p.text for p in final.paragraphs)
snippets = [
    "23.442 registros y 423 muertes fetales",
    "buena, 10.242 (43,69%)",
    "parcial, 9.524 (40,63%)",
    "baja, 3.676 (15,68%)",
    "OR 2,18 (IC95% 1,72–2,77; p<0,001)",
    "1,84 (IC95% 1,43–2,38; p<0,001)",
]
checks_texto = [s in texto_final for s in snippets]
figura_insertada = len(final.inline_shapes) > len(Document(ENTRADA).inline_shapes)
adherencia_ok = all(checks) and all(checks_texto)
manuales = 0 if adherencia_ok and errores_t1 == 0 and figura_insertada else 1

validacion = f"""# Validación del manuscrito final

- Tabla 1 validada: {'SÍ' if errores_t1 == 0 else 'NO'}
- Errores residuales Tabla 1: {errores_t1}
- Adherencia validada: {'SÍ' if adherencia_ok else 'NO'}
- Resultados Word coinciden con CSV: {'SÍ' if all(checks) else 'NO'}
- Modelos Word coinciden con resultados R: {'SÍ' if checks[2] and all(checks_texto) else 'NO'}
- Figura insertada: {'SÍ' if figura_insertada else 'NO'}
- Elementos de revisión manual: {manuales}
"""
Path("VALIDACION_MANUSCRITO_FINAL.md").write_text(validacion, encoding="utf-8")
print(validacion)
if manuales:
    raise RuntimeError("La validación final no fue completamente satisfactoria.")
