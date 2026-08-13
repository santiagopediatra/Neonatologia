#!/usr/bin/env python3

"""Consolida el manuscrito validado sin recalcular resultados."""

import csv
import shutil
import unicodedata
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

FUENTE = Path("manuscrito/Articulo_fetos_CORREGIDO.docx")
SALIDA = Path("manuscrito/Articulo_fetos_MAESTRO.docx")
INFORME = Path("VALIDACION_ARTICULO_MAESTRO.md")

REQUERIDOS = [
    Path("04_analisis_adherencia_prenatal.R"),
    Path("adherencia_auditoria.csv"), Path("adherencia_tabla_A1.csv"),
    Path("adherencia_tabla_A2.csv"), Path("adherencia_modelos_logisticos.csv"),
    Path("adherencia_estratificado.csv"), Path("adherencia_fig1.png"),
    Path("VALIDACION_MANUSCRITO_FINAL.md"),
    Path("INFORME_FINAL_AUDITORIA_Y_ADHERENCIA.md"), FUENTE,
]
faltantes = [str(p) for p in REQUERIDOS if not p.exists() or p.stat().st_size == 0]
if faltantes:
    raise RuntimeError("Insumos faltantes o vacíos: " + ", ".join(faltantes))

script = Path("04_analisis_adherencia_prenatal.R").read_text(encoding="utf-8")
for expresion in (
    "controles_observados =", "controles_esperados <- function(eg)",
    "adherencia_pct =", "categoria_adherencia = factor("
):
    if expresion not in script:
        raise RuntimeError(f"Definición no confirmada en el script: {expresion}")


def buscar(doc, inicio):
    encontrados = [p for p in doc.paragraphs if p.text.strip().startswith(inicio)]
    if len(encontrados) != 1:
        raise RuntimeError(f"Párrafo no encontrado inequívocamente: {inicio}")
    return encontrados[0]


def insertar_antes(parrafo, texto, estilo):
    xml = OxmlElement("w:p")
    parrafo._p.addprevious(xml)
    nuevo = Paragraph(xml, doc._body)
    pstyle = OxmlElement("w:pStyle")
    pstyle.set(qn("w:val"), estilo.replace(" ", ""))
    xml.get_or_add_pPr().append(pstyle)
    nuevo.add_run(texto)
    return nuevo


def insertar_despues(parrafo, texto, estilo="Body Text"):
    xml = OxmlElement("w:p")
    parrafo._p.addnext(xml)
    nuevo = Paragraph(xml, doc._body)
    pstyle = OxmlElement("w:pStyle")
    pstyle.set(qn("w:val"), estilo.replace(" ", ""))
    xml.get_or_add_pPr().append(pstyle)
    nuevo.add_run(texto)
    return nuevo


def reemplazar_texto(parrafo, texto):
    if not parrafo.runs:
        parrafo.add_run(texto)
        return
    parrafo.runs[0].text = texto
    for run in parrafo.runs[1:]:
        run.text = ""


shutil.copy2(FUENTE, SALIDA)
doc = Document(SALIDA)

# Separar la metodología adicional del párrafo general ya validado.
p_met = buscar(doc, "Se calcularon frecuencias absolutas")
marca_met = " Para el análisis adicional de adherencia"
if marca_met not in p_met.text:
    raise RuntimeError("No se encontró la metodología de adherencia previamente validada.")
general, detalle = p_met.text.split(marca_met, 1)
reemplazar_texto(p_met, general)
h_met = insertar_despues(p_met, "3.8.1 Adecuación o adherencia a los controles prenatales", "Heading 2")
p_det = insertar_despues(h_met, marca_met.strip() + detalle)
p_det.add_run(
    " El número observado procedió de NO_CONTROLES_NUM y la edad gestacional, expresada en semanas, de EDAD_GESTACIONAL_ANALITICA. La regla 1/2/3/6/8 aproximó el número acumulado esperado según la edad gestacional alcanzada y no se atribuyó como implementación textual exacta de una recomendación externa. Los valores superiores a 100% representaron más controles observados que los operacionalmente esperados; se conservaron sin truncamiento y todos permanecieron en la categoría buena. No se realizó imputación."
)

# Separar los resultados específicos y conservar las tablas/figura que ya siguen al párrafo.
p_res = buscar(doc, "Las asociaciones negativas de controles")
marca_res = " En el análisis operacional de adherencia"
if marca_res not in p_res.text:
    raise RuntimeError("No se encontraron los resultados de adherencia previamente validados.")
general_res, detalle_res = p_res.text.split(marca_res, 1)
reemplazar_texto(p_res, general_res)
h_res = insertar_despues(p_res, "4.9.1 Adherencia a controles prenatales", "Heading 2")
p_det_res = insertar_despues(h_res, marca_res.strip() + detalle_res)
p_det_res.add_run(
    " No hubo registros analíticos con edad gestacional <20 semanas. En el estrato 20–<28 semanas, los tamaños fueron reducidos y no sustentan un gradiente definitivo. Descriptivamente, la mortalidad fue mayor en baja adherencia en los estratos 28–<36 y ≥36 semanas."
)

# Consolidar discusión y limitaciones con las precisiones solicitadas.
p_disc = buscar(doc, "Las asociaciones negativas del modelo principal")
p_disc.add_run(
    " La adherencia parcial no mostró diferencias claras respecto de la buena adherencia; la baja adherencia mostró mayores odds, que permanecieron alejadas de la nulidad aunque con menor magnitud en la sensibilidad por edad gestacional. Esta persistencia no establece causalidad ni elimina explicaciones por oportunidad acumulada, temporalidad, selección o confusión residual."
)
p_lim = buscar(doc, "Las limitaciones incluyen diseño transversal")
p_lim.add_run(
    " Además, 15,75% de los registros analíticos tuvo adherencia >100%; estos valores se conservaron porque indican más controles observados que los operacionalmente esperados y no constituyen por sí mismos errores. El recuento de controles no mide la calidad, oportunidad ni contenido de la atención prenatal."
)

# Formato editorial de miles en la tabla estratificada, sin cambiar valores.
for tabla in doc.tables:
    encabezado = [c.text.strip() for c in tabla.rows[0].cells]
    if encabezado == ["Adherencia", "Desenlace", "n", "% columna", "Controles observados: media (DE)",
                      "Controles observados: mediana (RIC)", "Controles esperados: media (DE)", "EG: media (DE)"]:
        abreviados = ["Adherencia", "Desenlace", "n", "% columna", "Obs.: media (DE)",
                      "Obs.: mediana (RIC)", "Esperados: media (DE)", "EG: media (DE)"]
        for cell, texto in zip(tabla.rows[0].cells, abreviados):
            cell.text = texto
        encabezado = abreviados
    if encabezado == ["Estrato EG", "Adherencia", "N", "Muertes", "Mortalidad (%)", "IC95% inf", "IC95% sup"]:
        for row in tabla.rows[1:]:
            n = int(row.cells[2].text)
            row.cells[2].text = f"{n:,}".replace(",", ".")
            row.cells[4].text = row.cells[4].text.rstrip("%") + "%"
    if encabezado == ["Adherencia", "N", "Muertes", "Mortalidad (%)", "IC95% inf", "IC95% sup"]:
        for row in tabla.rows[1:]:
            row.cells[3].text = row.cells[3].text.rstrip("%") + "%"
    if encabezado and encabezado[0] in {"Adherencia", "Modelo", "Estrato EG"}:
        trPr = tabla.rows[0]._tr.get_or_add_trPr()
        tblHeader = OxmlElement("w:tblHeader")
        tblHeader.set(qn("w:val"), "true")
        trPr.append(tblHeader)
        for row in tabla.rows:
            rowPr = row._tr.get_or_add_trPr()
            cantSplit = OxmlElement("w:cantSplit")
            rowPr.append(cantSplit)
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(7.5)

doc.save(SALIDA)

# ------------------- Validación numérica tras reapertura -------------------
final = Document(SALIDA)
texto_parrafos = "\n".join(p.text for p in final.paragraphs)
texto_tablas = "\n".join(c.text for t in final.tables for row in t.rows for c in row.cells)
texto_total = texto_parrafos + "\n" + texto_tablas

valores_obligatorios = [
    "23.848", "23.642", "23.442", "423",
    "10.242", "9.524", "3.676", "3.693",
    "43,69%", "40,63%", "15,68%", "15,75%",
    "1,56%", "1,47%", "3,35%",
    "0,94", "0,75–1,18", "0,596",
    "2,18", "1,72–2,77", "<0,001",
    "1,08", "0,85–1,37", "0,528",
    "1,84", "1,43–2,38",
    "44", "26", "59,09%", "16", "7", "43,75%", "42", "61,90%",
    "681", "35", "5,14%", "239", "6,69%", "202", "24", "11,88%",
    "9.517", "99", "1,04%", "9.269", "117", "1,26%", "3.432", "73", "2,13%",
]
faltan_valores = [v for v in valores_obligatorios if v not in texto_total]

encabezados = [[c.text.strip() for c in t.rows[0].cells] for t in final.tables]
tablas_esperadas = [
    ["Variable", "Categoria", "MUERTE_FETAL_NO", "MUERTE_FETAL_SI"],
    ["Adherencia", "Desenlace", "n", "% columna", "Controles observados: media (DE)",
     "Controles observados: mediana (RIC)", "Controles esperados: media (DE)", "EG: media (DE)"],
    ["Adherencia", "N", "Muertes", "Mortalidad (%)", "IC95% inf", "IC95% sup"],
    ["Modelo", "Contraste", "OR", "IC95% inf", "IC95% sup", "p", "N", "Eventos"],
    ["Estrato EG", "Adherencia", "N", "Muertes", "Mortalidad (%)", "IC95% inf", "IC95% sup"],
]
tablas_esperadas[1] = ["Adherencia", "Desenlace", "n", "% columna", "Obs.: media (DE)",
                       "Obs.: mediana (RIC)", "Esperados: media (DE)", "EG: media (DE)"]
faltan_tablas = [x for x in tablas_esperadas if x not in encabezados]
figura_ok = len(final.inline_shapes) >= 5
subsecciones_ok = all(x in texto_parrafos for x in (
    "3.8.1 Adecuación o adherencia a los controles prenatales",
    "4.9.1 Adherencia a controles prenatales",
))

# Confirmar que Tabla 1 conserva los 56 valores validados.
with Path("tabla1_corregida.csv").open(encoding="utf-8-sig", newline="") as f:
    t1_csv = list(csv.DictReader(f))
t1 = next(t for t in final.tables if [c.text.strip() for c in t.rows[0].cells] == tablas_esperadas[0])

def norm(x):
    x = unicodedata.normalize("NFKD", x or "")
    return "".join(c for c in x if not unicodedata.combining(c)).casefold().strip()

esperados = {(norm(r["etiqueta"]), norm(r["categoria"]), r["desenlace"]): r["valor_formateado"] for r in t1_csv}
var = ""
errores_t1 = 0
for row in t1.rows[1:]:
    if row.cells[0].text.strip():
        var = row.cells[0].text.strip()
    cat = row.cells[1].text.strip()
    for col, des in ((2, "NO"), (3, "SI")):
        errores_t1 += row.cells[col].text.strip() != esperados[(norm(var), norm(cat), des)]

numerica_ok = not faltan_valores and not faltan_tablas and errores_t1 == 0 and subsecciones_ok
if not numerica_ok or not figura_ok:
    raise RuntimeError(
        f"Validación fallida; valores ausentes={faltan_valores}; tablas ausentes={len(faltan_tablas)}; "
        f"errores Tabla 1={errores_t1}; subsecciones={subsecciones_ok}; figura={figura_ok}"
    )

informe = f"""# Validación del artículo maestro

- Documento fuente: `manuscrito/Articulo_fetos_CORREGIDO.docx`
- Documento final: `manuscrito/Articulo_fetos_MAESTRO.docx`
- Secciones actualizadas: Métodos; Resultados; Discusión; Limitaciones.
- Resumen: conservó la actualización validada de adherencia del documento fuente.
- Abstract: no existe como sección en el documento fuente; no se inventó una traducción.
- Subsecciones creadas: `3.8.1 Adecuación o adherencia a los controles prenatales`; `4.9.1 Adherencia a controles prenatales`.
- Tablas conservadas e integradas: Tabla 1 y tablas A1–A4.
- Figura insertada y conservada: Figura A1.
- N analítico de adherencia: 23.442; eventos: 423.
- Comprobación de resultados numéricos: OK.
- Errores residuales Tabla 1: {errores_t1}.
- Discrepancias encontradas: ninguna numérica.
- Discrepancias corregidas en esta consolidación: separación estructural de metodología y resultados; aclaración de valores >100% y estratificación.
- Numeración A1–A4: provisional y sin colisiones; la numeración editorial final requiere revisión.
- Revisión humana pendiente: decisión editorial sobre numeración definitiva y eventual incorporación de Abstract; aprobación o referencia externa de la regla operacional 1/2/3/6/8.
"""
INFORME.write_text(informe, encoding="utf-8")

print("Confirmado en el script: controles_observados, controles_esperados, adherencia_pct y categoria_adherencia")
print("Validación numérica: OK")
print(f"Tablas totales: {len(final.tables)}; figuras totales: {len(final.inline_shapes)}")
print(f"Documento maestro generado: {SALIDA}")
